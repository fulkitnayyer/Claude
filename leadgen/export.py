"""Rebuild the Batch 8 workbook from SQLite. Called after every accepted clinic.

Uses the master tracker as a template so the styling, Weights sheet and Status/Ads
dropdowns match. The file is written to a temp path then atomically renamed, so a
crash mid-write never leaves a corrupt xlsx behind.
"""
import copy
import datetime as dt
import os

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

from config import HEADERS, MASTER_XLSX, OUT_XLSX

STATUS_LIST = '"Not started,Message drafted,Sent,Follow-up 1,Follow-up 2,Replied - interested,Replied - no,No response,Do not contact"'
ADS_LIST = '"Running,Lapsed,Not running,Not checked yet"'
EXTRA_WIDTHS = {26: 18, 27: 38, 28: 40}


def build(leads, drops=None):
    wb = openpyxl.load_workbook(MASTER_XLSX)
    ws = wb["Delhi Prospects"]
    head_style = ws.cell(1, 2)
    body_style = ws.cell(2, 2)

    ws.delete_rows(2, ws.max_row)
    ws.data_validations.dataValidation = []
    ws.conditional_formatting = type(ws.conditional_formatting)()

    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, c, h)
        if c > 25:
            cell._style = copy.copy(head_style._style)
            ws.column_dimensions[cell.column_letter].width = EXTRA_WIDTHS[c]

    last = len(leads) + 1
    for i, lead in enumerate(leads, 2):
        values = [lead.get(h) for h in HEADERS]
        values[16] = f"=N{i}*Weights!$B$2+O{i}*Weights!$B$3+P{i}*Weights!$B$4"
        values[17] = f"=RANK(Q{i},$Q$2:$Q${last})"
        for c, v in enumerate(values, 1):
            cell = ws.cell(i, c, v)
            cell._style = copy.copy(body_style._style)

    if leads:
        for rng, formula in ((f"S2:S{last}", STATUS_LIST), (f"W2:W{last}", ADS_LIST)):
            dv = DataValidation(type="list", formula1=formula, allow_blank=True)
            dv.add(rng)
            ws.add_data_validation(dv)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(HEADERS)).column_letter}{last}"

    _about(wb["About"], leads, drops or [])
    if "Reserves" in wb.sheetnames:
        del wb["Reserves"]

    tmp = OUT_XLSX.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, OUT_XLSX)


def _about(ws, leads, drops):
    ws.delete_rows(1, ws.max_row)
    n_ads = sum(1 for l in leads if str(l.get("Practo Ads (Sponsored/Prime)", "")).startswith("Yes"))
    n_ver = sum(1 for l in leads if str(l.get("Phone Confidence", "")).startswith("Verified"))
    lines = [
        "HealthGrow – Delhi Dental Prospect Tracker",
        f"Batch 8 – Practo-sourced build – last updated {dt.datetime.now():%d %b %Y %H:%M}",
        "",
        "Scope",
        "Independent Delhi dental clinics (1–4 branches) listed on Practo, excluding the 73 clinics already "
        "in the master tracker (matched by phone, website domain and clinic name). Numbering continues from #74.",
        "",
        "Progress",
        f"{len(leads)} clinics accepted so far. {n_ver} have a phone verified on the clinic's own website; "
        f"{n_ads} were seen running Practo ads (Sponsored slot or Prime badge).",
        "",
        "Source & verification",
        "Each clinic comes from its Practo clinic profile (name, doctors, qualifications, address, fees, reviews). "
        "Where the profile links an own website, that site was fetched to pull a direct phone, email and WhatsApp "
        "link. Phone Confidence: 'Verified: clinic's own website' = seen on own site; 'Directory (Practo)' = only "
        "seen on Practo, which is usually a call-forwarding number – confirm before WhatsApp outreach.",
        "",
        "Practo Ads column",
        "'Yes – Sponsored' means the clinic appeared in a Sponsored slot on a Practo listing page during collection "
        "(locality and page are in the evidence column). 'Yes – Prime' means the profile carries Practo's paid Prime "
        "badge. Either one shows the clinic already pays for patient acquisition → warm 'is your spend working?' "
        "angle. Google Ads columns are intentionally left as 'Not checked yet'.",
        "",
        "Scoring",
        "Same formula as the master: Score = Fit×0.4 + Signal×0.4 + Reachability×0.2 (editable on Weights). "
        "Fit = independence + core services (implants, aligners, RCT, cosmetic) + MDS specialists. "
        "Signal = Practo ad spend, review volume, multi-dentist bench. Reachability = verified mobile, WhatsApp, email.",
        "",
        "Anti-fabrication note",
        "Nothing is invented. Anything not found on Practo or the clinic's own site is marked 'Not found'.",
    ]
    if drops:
        lines += ["", "Dropped candidates (reason – count)"] + [f"{r} – {n}" for r, n in drops]
    for i, t in enumerate(lines, 1):
        ws.cell(i, 1, t)
    ws.column_dimensions["A"].width = 140
